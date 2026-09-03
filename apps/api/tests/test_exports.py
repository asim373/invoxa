import csv
import io
import re
import uuid
import zipfile
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from apps.api.app.auth import create_access_token
from apps.api.app.database import get_db
from apps.api.app.main import app
from apps.api.app.models import (
    Document,
    DocumentStatus,
    InvoiceExtraction,
    InvoiceLineItem,
    User,
)


class FakeResult:
    def __init__(self, rows) -> None:
        self.rows = rows

    def yield_per(self, size: int):
        assert size > 0
        return self

    def all(self):
        raise AssertionError("Exports must iterate in batches, not call all().")

    def __iter__(self):
        return iter(self.rows)


class FakeExportDatabase:
    def __init__(self, users, documents, extractions, line_items) -> None:
        self.users = users
        self.documents = documents
        self.extractions = extractions
        self.line_items = line_items

    def get(self, model, identifier):
        if model is User:
            return next((user for user in self.users if user.id == identifier), None)
        return None

    def scalars(self, statement):
        ids, owner_id = self._ids_and_owner(statement)
        return FakeResult(
            [
                document.id
                for document in self.documents
                if document.id in ids
                and document.owner_id == owner_id
                and document.owner_id is not None
            ]
        )

    def execute(self, statement):
        ids, owner_id = self._ids_and_owner(statement)
        documents = sorted(
            (
                document
                for document in self.documents
                if document.id in ids
                and document.owner_id == owner_id
                and document.owner_id is not None
            ),
            key=lambda document: (document.created_at, document.id),
        )
        selected_names = [column["name"] for column in statement.column_descriptions]
        if "position" not in selected_names:
            return FakeResult(
                [(document, self.extractions.get(document.id)) for document in documents]
            )

        rows = []
        for document in documents:
            extraction = self.extractions.get(document.id)
            if extraction is None:
                continue
            for item in sorted(
                (
                    line_item
                    for line_item in self.line_items
                    if line_item.invoice_document_id == document.id
                ),
                key=lambda line_item: line_item.position,
            ):
                rows.append(
                    (
                        document.id,
                        extraction.invoice_number,
                        item.position,
                        item.description,
                        item.quantity,
                        item.unit_price,
                        item.line_total,
                    )
                )
        return FakeResult(rows)

    @staticmethod
    def _ids_and_owner(statement):
        parameters = statement.compile().params
        ids = next(value for value in parameters.values() if isinstance(value, (list, tuple)))
        owner_id = next(value for value in parameters.values() if isinstance(value, uuid.UUID))
        return set(ids), owner_id


def make_document(owner_id, created_at, filename, status=DocumentStatus.PROCESSING):
    return Document(
        id=uuid.uuid4(),
        owner_id=owner_id,
        original_filename=filename,
        mime_type="application/pdf",
        file_size=1,
        checksum="a" * 64,
        storage_location="internal.pdf",
        extracted_text="TOP SECRET OCR TEXT",
        error_details="INTERNAL ERROR DETAIL",
        status=status,
        created_at=created_at,
        updated_at=created_at,
    )


def make_extraction(document_id, **overrides):
    values = {
        "document_id": document_id,
        "invoice_number": "INV-100",
        "invoice_date": date(2026, 8, 24),
        "vendor_name": "Acme, Inc.",
        "customer_name": "Customer",
        "currency": "USD",
        "subtotal": Decimal("1200.50"),
        "tax": Decimal("240.10"),
        "total": Decimal("1440.60"),
        "is_valid": True,
        "validation_error": None,
    }
    values.update(overrides)
    return InvoiceExtraction(**values)


def make_line_item(document_id, position, description, quantity="2.0000", price="10.00"):
    quantity_decimal = Decimal(quantity)
    price_decimal = Decimal(price)
    return InvoiceLineItem(
        id=uuid.uuid4(),
        invoice_document_id=document_id,
        position=position,
        description=description,
        quantity=quantity_decimal,
        unit_price=price_decimal,
        line_total=(quantity_decimal * price_decimal).quantize(Decimal("0.01")),
    )


@pytest.fixture
def export_data():
    now = datetime(2026, 8, 24, tzinfo=UTC)
    owner = User(
        id=uuid.uuid4(),
        email="owner@example.com",
        password_hash="unused",
        is_active=True,
    )
    other = User(
        id=uuid.uuid4(),
        email="other@example.com",
        password_hash="unused",
        is_active=True,
    )
    first = make_document(owner.id, now, "first.pdf")
    second = make_document(owner.id, now + timedelta(seconds=1), "second.pdf")
    other_document = make_document(other.id, now + timedelta(seconds=2), "other.pdf")
    legacy_document = make_document(None, now + timedelta(seconds=3), "legacy.pdf")
    extractions = {
        first.id: make_extraction(first.id),
        second.id: make_extraction(
            second.id,
            invoice_number="INV-200",
            is_valid=False,
            validation_error="Subtotal mismatch",
        ),
        other_document.id: make_extraction(other_document.id),
    }
    line_items = [
        make_line_item(first.id, 1, "Second item", quantity="3.0000", price="5.50"),
        make_line_item(first.id, 0, "First item"),
        make_line_item(second.id, 0, "Second invoice item"),
    ]
    database = FakeExportDatabase(
        [owner, other],
        [second, legacy_document, first, other_document],
        extractions,
        line_items,
    )
    return database, owner, other, first, second, other_document, legacy_document


@pytest.fixture
def client(export_data):
    database = export_data[0]
    app.dependency_overrides[get_db] = lambda: database
    yield TestClient(app)
    app.dependency_overrides.clear()


def headers_for(user):
    return {"Authorization": f"Bearer {create_access_token(user.id)}"}


def read_zip(response):
    archive = zipfile.ZipFile(io.BytesIO(response.content))
    assert archive.namelist() == ["invoices.csv", "line_items.csv"]
    return {
        name: list(
            csv.reader(io.TextIOWrapper(archive.open(name), encoding="utf-8-sig", newline=""))
        )
        for name in archive.namelist()
    }


# ===================================================================
# CSV export tests
# ===================================================================


def test_single_owned_document_exports_exact_zip_and_headers(client, export_data) -> None:
    _, owner, _, first, *_ = export_data
    response = client.get(f"/documents/{first.id}/exports/csv", headers=headers_for(owner))

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/zip")
    assert re.fullmatch(
        r'attachment; filename="invoice-export-\d{8}T\d{6}Z\.zip"',
        response.headers["content-disposition"],
    )
    files = read_zip(response)
    assert files["invoices.csv"][0] == [
        "document_id",
        "filename",
        "document_status",
        "invoice_number",
        "invoice_date",
        "vendor_name",
        "customer_name",
        "currency",
        "subtotal",
        "tax",
        "total",
        "is_valid",
        "validation_error",
    ]
    assert files["line_items.csv"][0] == [
        "document_id",
        "invoice_number",
        "line_item_index",
        "description",
        "quantity",
        "unit_price",
        "line_total",
    ]
    assert len(files["invoices.csv"]) == 2
    assert len(files["line_items.csv"]) == 3


def test_bulk_export_is_deterministic_and_preserves_invoice_and_line_item_data(
    client, export_data
) -> None:
    _, owner, _, first, second, *_ = export_data
    response = client.post(
        "/exports/csv",
        headers=headers_for(owner),
        json={"document_ids": [str(second.id), str(first.id), str(first.id)]},
    )

    files = read_zip(response)
    invoices = files["invoices.csv"]
    line_items = files["line_items.csv"]
    assert [row[0] for row in invoices[1:]] == [str(first.id), str(second.id)]
    assert invoices[1][8:13] == ["1200.50", "240.10", "1440.60", "true", ""]
    assert invoices[2][3:13] == [
        "INV-200",
        "2026-08-24",
        "Acme, Inc.",
        "Customer",
        "USD",
        "1200.50",
        "240.10",
        "1440.60",
        "false",
        "Subtotal mismatch",
    ]
    assert [(row[0], row[2], row[3]) for row in line_items[1:]] == [
        (str(first.id), "0", "First item"),
        (str(first.id), "1", "Second item"),
        (str(second.id), "0", "Second invoice item"),
    ]
    assert line_items[1][4:] == ["2.0000", "10.00", "20.00"]


def test_missing_extraction_is_retained_as_blank_invoice_row(client, export_data) -> None:
    database, owner, _, _, _, _, _ = export_data
    document = make_document(owner.id, datetime(2026, 8, 25, tzinfo=UTC), "unprocessed.pdf")
    database.documents.append(document)

    response = client.get(f"/documents/{document.id}/exports/csv", headers=headers_for(owner))

    files = read_zip(response)
    assert files["invoices.csv"][1] == [
        str(document.id),
        "unprocessed.pdf",
        "processing",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
    ]
    assert files["line_items.csv"] == [
        [
            "document_id",
            "invoice_number",
            "line_item_index",
            "description",
            "quantity",
            "unit_price",
            "line_total",
        ]
    ]


def test_csv_text_is_utf8_quoted_and_formula_safe(client, export_data) -> None:
    database, owner, _, first, _, _, _ = export_data
    extraction = database.extractions[first.id]
    extraction.invoice_number = " =INJECT()"
    extraction.vendor_name = 'موجود, "quoted"\nnext line'
    extraction.customer_name = "@formula"
    extraction.validation_error = "\tformula"
    extraction.is_valid = False
    database.line_items.append(make_line_item(first.id, 2, "-description"))
    response = client.get(f"/documents/{first.id}/exports/csv", headers=headers_for(owner))

    files = read_zip(response)
    invoice = files["invoices.csv"][1]
    assert invoice[3] == "' =INJECT()"
    assert invoice[5] == 'موجود, "quoted"\nnext line'
    assert invoice[6] == "'@formula"
    assert invoice[12] == "'\tformula"
    assert files["line_items.csv"][-1][3] == "'-description"


@pytest.mark.parametrize("kind", ["other", "legacy", "missing"])
def test_export_denies_unowned_or_unowned_documents_without_partial_zip(
    client, export_data, kind
) -> None:
    _, owner, _, first, _, other_document, legacy_document = export_data
    denied_id = {
        "other": other_document.id,
        "legacy": legacy_document.id,
        "missing": uuid.uuid4(),
    }[kind]
    response = client.post(
        "/exports/csv",
        headers=headers_for(owner),
        json={"document_ids": [str(first.id), str(denied_id)]},
    )

    assert response.status_code == 404
    assert response.headers["content-type"].startswith("application/json")
    assert response.json() == {"detail": "Document not found."}


def test_export_requires_active_authenticated_user(client, export_data) -> None:
    database, owner, _, first, *_ = export_data
    unauthenticated = client.get(f"/documents/{first.id}/exports/csv")
    owner.is_active = False
    inactive = client.get(f"/documents/{first.id}/exports/csv", headers=headers_for(owner))

    assert unauthenticated.status_code == 401
    assert inactive.status_code == 401
    assert database.documents


def test_empty_bulk_export_is_rejected(client, export_data) -> None:
    _, owner, *_ = export_data
    response = client.post("/exports/csv", headers=headers_for(owner), json={"document_ids": []})

    assert response.status_code == 422


def test_export_does_not_expose_internal_document_fields(client, export_data) -> None:
    _, owner, _, first, *_ = export_data
    response = client.get(f"/documents/{first.id}/exports/csv", headers=headers_for(owner))

    assert b"internal.pdf" not in response.content
    assert b"TOP SECRET OCR TEXT" not in response.content
    assert b"INTERNAL ERROR DETAIL" not in response.content
    assert b"checksum" not in response.content


# ===================================================================
# XLSX export tests
# ===================================================================

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def read_xlsx(response):
    return load_workbook(io.BytesIO(response.content))


def sheet_rows(sheet):
    """Return all rows as lists of cell values (including header)."""
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def test_xlsx_single_owned_document_exports_valid_workbook(client, export_data) -> None:
    _, owner, _, first, *_ = export_data
    response = client.get(f"/documents/{first.id}/exports/xlsx", headers=headers_for(owner))

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(XLSX_MEDIA_TYPE)
    assert re.fullmatch(
        r'attachment; filename="invoice-export-\d{8}T\d{6}Z\.xlsx"',
        response.headers["content-disposition"],
    )


def test_xlsx_has_exactly_two_worksheets(client, export_data) -> None:
    _, owner, _, first, *_ = export_data
    response = client.get(f"/documents/{first.id}/exports/xlsx", headers=headers_for(owner))
    workbook = read_xlsx(response)

    assert workbook.sheetnames == ["Invoices", "Line Items"]


def test_xlsx_invoices_sheet_has_exact_headers(client, export_data) -> None:
    _, owner, _, first, *_ = export_data
    response = client.get(f"/documents/{first.id}/exports/xlsx", headers=headers_for(owner))
    workbook = read_xlsx(response)
    invoices = sheet_rows(workbook["Invoices"])

    assert invoices[0] == [
        "document_id",
        "filename",
        "document_status",
        "invoice_number",
        "invoice_date",
        "vendor_name",
        "customer_name",
        "currency",
        "subtotal",
        "tax",
        "total",
        "is_valid",
        "validation_error",
    ]


def test_xlsx_line_items_sheet_has_exact_headers(client, export_data) -> None:
    _, owner, _, first, *_ = export_data
    response = client.get(f"/documents/{first.id}/exports/xlsx", headers=headers_for(owner))
    workbook = read_xlsx(response)
    line_items = sheet_rows(workbook["Line Items"])

    assert line_items[0] == [
        "document_id",
        "invoice_number",
        "line_item_index",
        "description",
        "quantity",
        "unit_price",
        "line_total",
    ]


def test_xlsx_single_document_has_one_invoice_row(client, export_data) -> None:
    _, owner, _, first, *_ = export_data
    response = client.get(f"/documents/{first.id}/exports/xlsx", headers=headers_for(owner))
    workbook = read_xlsx(response)
    invoices = sheet_rows(workbook["Invoices"])

    assert len(invoices) == 2  # header + 1 data row


def test_xlsx_single_document_line_items_appear_exactly_once(client, export_data) -> None:
    _, owner, _, first, *_ = export_data
    response = client.get(f"/documents/{first.id}/exports/xlsx", headers=headers_for(owner))
    workbook = read_xlsx(response)
    line_items = sheet_rows(workbook["Line Items"])

    # first has 2 line items
    assert len(line_items) == 3  # header + 2 data rows


def test_xlsx_bulk_export_is_deterministic_and_preserves_data(client, export_data) -> None:
    _, owner, _, first, second, *_ = export_data
    response = client.post(
        "/exports/xlsx",
        headers=headers_for(owner),
        json={"document_ids": [str(second.id), str(first.id), str(first.id)]},
    )

    assert response.status_code == 200
    workbook = read_xlsx(response)
    invoices = sheet_rows(workbook["Invoices"])
    line_items = sheet_rows(workbook["Line Items"])

    # Document ordering: first (earlier created_at) then second
    assert [row[0] for row in invoices[1:]] == [str(first.id), str(second.id)]

    # Invoice data preserved
    assert invoices[1][8:11] == ["1200.50", "240.10", "1440.60"]
    assert invoices[1][11] is True
    assert invoices[2][3] == "INV-200"
    assert invoices[2][11] is False
    assert invoices[2][12] == "Subtotal mismatch"

    # Line items: first's items (position 0, 1) then second's item (position 0)
    assert [(row[0], row[2], row[3]) for row in line_items[1:]] == [
        (str(first.id), 0, "First item"),
        (str(first.id), 1, "Second item"),
        (str(second.id), 0, "Second invoice item"),
    ]


def test_xlsx_position_ordering(client, export_data) -> None:
    _, owner, _, first, *_ = export_data
    response = client.get(f"/documents/{first.id}/exports/xlsx", headers=headers_for(owner))
    workbook = read_xlsx(response)
    line_items = sheet_rows(workbook["Line Items"])

    positions = [row[2] for row in line_items[1:]]
    assert positions == [0, 1]


def test_xlsx_zero_line_item_invoice(client, export_data) -> None:
    database, owner, _, _, _, _, _ = export_data
    document = make_document(owner.id, datetime(2026, 8, 25, tzinfo=UTC), "no-items.pdf")
    database.documents.append(document)
    database.extractions[document.id] = make_extraction(document.id)

    response = client.get(f"/documents/{document.id}/exports/xlsx", headers=headers_for(owner))
    workbook = read_xlsx(response)
    invoices = sheet_rows(workbook["Invoices"])
    line_items = sheet_rows(workbook["Line Items"])

    assert len(invoices) == 2  # header + 1 invoice row
    assert len(line_items) == 1  # header only, no line items


def test_xlsx_missing_extraction_is_blank_row(client, export_data) -> None:
    database, owner, _, _, _, _, _ = export_data
    document = make_document(owner.id, datetime(2026, 8, 25, tzinfo=UTC), "unprocessed.pdf")
    database.documents.append(document)

    response = client.get(f"/documents/{document.id}/exports/xlsx", headers=headers_for(owner))
    workbook = read_xlsx(response)
    invoices = sheet_rows(workbook["Invoices"])

    row = invoices[1]
    assert row[0] == str(document.id)
    assert row[1] == "unprocessed.pdf"
    assert row[2] == "processing"
    # All extraction fields should be None
    assert row[3:] == [None, None, None, None, None, None, None, None, None, None]

    line_items = sheet_rows(workbook["Line Items"])
    assert len(line_items) == 1  # header only


def test_xlsx_invalid_extraction(client, export_data) -> None:
    _, owner, _, _, second, *_ = export_data
    response = client.get(f"/documents/{second.id}/exports/xlsx", headers=headers_for(owner))
    workbook = read_xlsx(response)
    invoices = sheet_rows(workbook["Invoices"])

    row = invoices[1]
    assert row[11] is False
    assert row[12] == "Subtotal mismatch"


def test_xlsx_decimal_precision(client, export_data) -> None:
    _, owner, _, first, *_ = export_data
    response = client.get(f"/documents/{first.id}/exports/xlsx", headers=headers_for(owner))
    workbook = read_xlsx(response)
    invoices = sheet_rows(workbook["Invoices"])
    line_items = sheet_rows(workbook["Line Items"])

    # Money fields are strings with exactly 2 decimal places
    assert invoices[1][8] == "1200.50"
    assert invoices[1][9] == "240.10"
    assert invoices[1][10] == "1440.60"

    # Quantity is string with exactly 4 decimal places
    assert line_items[1][4] == "2.0000"
    # Unit price and line total are money (2 decimal)
    assert line_items[1][5] == "10.00"
    assert line_items[1][6] == "20.00"


def test_xlsx_date_formatting(client, export_data) -> None:
    _, owner, _, first, *_ = export_data
    response = client.get(f"/documents/{first.id}/exports/xlsx", headers=headers_for(owner))
    workbook = read_xlsx(response)

    # Access the actual cell (not just values) to check format
    invoice_sheet = workbook["Invoices"]
    date_cell = invoice_sheet.cell(row=2, column=5)  # invoice_date column
    assert isinstance(date_cell.value, datetime)  # openpyxl reads dates as datetimes
    assert date_cell.value.date() == date(2026, 8, 24)
    assert date_cell.number_format == "YYYY-MM-DD"


def test_xlsx_currency_preserved(client, export_data) -> None:
    _, owner, _, first, *_ = export_data
    response = client.get(f"/documents/{first.id}/exports/xlsx", headers=headers_for(owner))
    workbook = read_xlsx(response)
    invoices = sheet_rows(workbook["Invoices"])

    assert invoices[1][7] == "USD"


def test_xlsx_null_values(client, export_data) -> None:
    database, owner, _, _, _, _, _ = export_data
    document = make_document(owner.id, datetime(2026, 8, 25, tzinfo=UTC), "nulls.pdf")
    database.documents.append(document)
    database.extractions[document.id] = make_extraction(
        document.id,
        invoice_number=None,
        invoice_date=None,
        vendor_name=None,
        customer_name=None,
        currency=None,
        subtotal=None,
        tax=None,
        total=None,
    )
    response = client.get(f"/documents/{document.id}/exports/xlsx", headers=headers_for(owner))
    workbook = read_xlsx(response)
    invoices = sheet_rows(workbook["Invoices"])

    row = invoices[1]
    assert row[3] is None  # invoice_number
    assert row[4] is None  # invoice_date
    assert row[5] is None  # vendor_name
    assert row[6] is None  # customer_name
    assert row[7] is None  # currency
    assert row[8] is None  # subtotal
    assert row[9] is None  # tax
    assert row[10] is None  # total


def test_xlsx_unicode_quotes_and_multiline(client, export_data) -> None:
    database, owner, _, first, _, _, _ = export_data
    extraction = database.extractions[first.id]
    extraction.vendor_name = 'موجود, "quoted"\nnext line'
    extraction.customer_name = "日本語テスト"

    response = client.get(f"/documents/{first.id}/exports/xlsx", headers=headers_for(owner))
    workbook = read_xlsx(response)
    invoices = sheet_rows(workbook["Invoices"])

    assert invoices[1][5] == 'موجود, "quoted"\nnext line'
    assert invoices[1][6] == "日本語テスト"


def test_xlsx_formula_injection_payloads(client, export_data) -> None:
    database, owner, _, first, _, _, _ = export_data
    extraction = database.extractions[first.id]
    extraction.invoice_number = " =INJECT()"
    extraction.vendor_name = "+cmd"
    extraction.customer_name = "@formula"
    extraction.validation_error = "\tformula"
    extraction.is_valid = False
    database.line_items.append(make_line_item(first.id, 2, "-description"))

    response = client.get(f"/documents/{first.id}/exports/xlsx", headers=headers_for(owner))
    workbook = read_xlsx(response)
    invoices = sheet_rows(workbook["Invoices"])
    line_items = sheet_rows(workbook["Line Items"])

    invoice = invoices[1]
    assert invoice[3] == "' =INJECT()"
    assert invoice[5] == "'+cmd"
    assert invoice[6] == "'@formula"
    assert invoice[12] == "'\tformula"
    assert line_items[-1][3] == "'-description"


@pytest.mark.parametrize("kind", ["other", "legacy", "missing"])
def test_xlsx_export_denies_unowned_documents(client, export_data, kind) -> None:
    _, owner, _, first, _, other_document, legacy_document = export_data
    denied_id = {
        "other": other_document.id,
        "legacy": legacy_document.id,
        "missing": uuid.uuid4(),
    }[kind]
    response = client.post(
        "/exports/xlsx",
        headers=headers_for(owner),
        json={"document_ids": [str(first.id), str(denied_id)]},
    )

    assert response.status_code == 404
    assert response.headers["content-type"].startswith("application/json")
    assert response.json() == {"detail": "Document not found."}


def test_xlsx_export_requires_active_authenticated_user(client, export_data) -> None:
    database, owner, _, first, *_ = export_data
    unauthenticated = client.get(f"/documents/{first.id}/exports/xlsx")
    owner.is_active = False
    inactive = client.get(f"/documents/{first.id}/exports/xlsx", headers=headers_for(owner))

    assert unauthenticated.status_code == 401
    assert inactive.status_code == 401
    assert database.documents


def test_xlsx_empty_bulk_export_is_rejected(client, export_data) -> None:
    _, owner, *_ = export_data
    response = client.post("/exports/xlsx", headers=headers_for(owner), json={"document_ids": []})

    assert response.status_code == 422


def test_xlsx_server_generated_safe_filename(client, export_data) -> None:
    _, owner, _, first, *_ = export_data
    response = client.get(f"/documents/{first.id}/exports/xlsx", headers=headers_for(owner))

    disposition = response.headers["content-disposition"]
    assert re.fullmatch(
        r'attachment; filename="invoice-export-\d{8}T\d{6}Z\.xlsx"',
        disposition,
    )
    # Ensure user-uploaded filename is NOT in the download filename
    assert "first.pdf" not in disposition


def test_xlsx_export_does_not_expose_internal_document_fields(client, export_data) -> None:
    _, owner, _, first, *_ = export_data
    response = client.get(f"/documents/{first.id}/exports/xlsx", headers=headers_for(owner))
    workbook = read_xlsx(response)

    # Check all sheet values for internal field leaks
    for sheet_name in workbook.sheetnames:
        for row in workbook[sheet_name].iter_rows(values_only=True):
            for cell_value in row:
                if cell_value is None:
                    continue
                cell_str = str(cell_value)
                assert "internal.pdf" not in cell_str
                assert "TOP SECRET OCR TEXT" not in cell_str
                assert "INTERNAL ERROR DETAIL" not in cell_str


def test_xlsx_boolean_values_are_native(client, export_data) -> None:
    _, owner, _, first, second, *_ = export_data
    response = client.post(
        "/exports/xlsx",
        headers=headers_for(owner),
        json={"document_ids": [str(first.id), str(second.id)]},
    )
    workbook = read_xlsx(response)
    invoices = sheet_rows(workbook["Invoices"])

    # first has is_valid=True, second has is_valid=False
    assert invoices[1][11] is True
    assert invoices[2][11] is False


def test_xlsx_freeze_panes_and_autofilter(client, export_data) -> None:
    _, owner, _, first, *_ = export_data
    response = client.get(f"/documents/{first.id}/exports/xlsx", headers=headers_for(owner))
    workbook = read_xlsx(response)

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref is not None
