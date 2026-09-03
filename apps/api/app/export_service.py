import csv
import io
import tempfile
import uuid
import zipfile
from collections.abc import Iterable
from datetime import date
from decimal import Decimal
from typing import Any, BinaryIO, cast

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from apps.api.app.models import Document, InvoiceExtraction, InvoiceLineItem

INVOICE_HEADERS = (
    "document_id",
    "filename",
    "document_status",
    "invoice_number",
    "invoice_date",
    "vendor_name",
    "customer_name",
    "currency",
    "subtotal",
    "tax",
    "total",
    "is_valid",
    "validation_error",
)
LINE_ITEM_HEADERS = (
    "document_id",
    "invoice_number",
    "line_item_index",
    "description",
    "quantity",
    "unit_price",
    "line_total",
)
QUERY_BATCH_SIZE = 500
SPOOL_MEMORY_LIMIT = 8 * 1024 * 1024


class ExportDocumentsNotFoundError(ValueError):
    pass


# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------


def sanitize_cell_text(value: str | None) -> str:
    """Sanitize text for safe use in spreadsheet cells (CSV and XLSX).

    Prevents formula injection by prefixing dangerous leading characters
    with a single quote.
    """
    if value is None:
        return ""
    if value[:1] in {"\t", "\r", "\n"} or value.lstrip().startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


# Backward-compatible alias so existing CSV code keeps working unchanged.
sanitize_csv_text = sanitize_cell_text


def serialize_money(value: Decimal | None) -> str:
    return "" if value is None else format(value, ".2f")


def serialize_quantity(value: Decimal | None) -> str:
    return "" if value is None else format(value, ".4f")


def serialize_date(value: date | None) -> str:
    return "" if value is None else value.isoformat()


# ---------------------------------------------------------------------------
# CSV export
# ---------------------------------------------------------------------------


def create_csv_export(
    database: Session, owner_id: uuid.UUID, requested_document_ids: Iterable[uuid.UUID]
) -> BinaryIO:
    document_ids = tuple(dict.fromkeys(requested_document_ids))
    if not document_ids:
        raise ValueError("At least one document ID is required.")

    _verify_owned_documents(database, owner_id, document_ids)
    export_file = tempfile.SpooledTemporaryFile(max_size=SPOOL_MEMORY_LIMIT, mode="w+b")
    try:
        with zipfile.ZipFile(export_file, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
            _write_invoices_csv(archive, database, owner_id, document_ids)
            _write_line_items_csv(archive, database, owner_id, document_ids)
        export_file.seek(0)
        return cast(BinaryIO, export_file)
    except Exception:
        export_file.close()
        raise


def _verify_owned_documents(
    database: Session, owner_id: uuid.UUID, document_ids: tuple[uuid.UUID, ...]
) -> None:
    statement = select(Document.id).where(
        Document.id.in_(document_ids),
        Document.owner_id == owner_id,
        Document.owner_id.is_not(None),
    )
    found_ids = set(database.scalars(statement).yield_per(QUERY_BATCH_SIZE))
    if found_ids != set(document_ids):
        raise ExportDocumentsNotFoundError


def _write_invoices_csv(
    archive: zipfile.ZipFile,
    database: Session,
    owner_id: uuid.UUID,
    document_ids: tuple[uuid.UUID, ...],
) -> None:
    statement = (
        select(Document, InvoiceExtraction)
        .outerjoin(InvoiceExtraction, InvoiceExtraction.document_id == Document.id)
        .where(
            Document.id.in_(document_ids),
            Document.owner_id == owner_id,
            Document.owner_id.is_not(None),
        )
        .order_by(Document.created_at.asc(), Document.id.asc())
    )
    with archive.open("invoices.csv", mode="w") as binary_file:
        text_file = io.TextIOWrapper(binary_file, encoding="utf-8-sig", newline="")
        try:
            writer = csv.writer(text_file, lineterminator="\r\n")
            writer.writerow(INVOICE_HEADERS)
            for document, extraction in database.execute(statement).yield_per(QUERY_BATCH_SIZE):
                writer.writerow(_invoice_row(document, extraction))
            text_file.flush()
        finally:
            text_file.detach()


def _write_line_items_csv(
    archive: zipfile.ZipFile,
    database: Session,
    owner_id: uuid.UUID,
    document_ids: tuple[uuid.UUID, ...],
) -> None:
    statement = (
        select(
            Document.id,
            InvoiceExtraction.invoice_number,
            InvoiceLineItem.position,
            InvoiceLineItem.description,
            InvoiceLineItem.quantity,
            InvoiceLineItem.unit_price,
            InvoiceLineItem.line_total,
        )
        .select_from(Document)
        .join(InvoiceExtraction, InvoiceExtraction.document_id == Document.id)
        .join(
            InvoiceLineItem,
            InvoiceLineItem.invoice_document_id == InvoiceExtraction.document_id,
        )
        .where(
            Document.id.in_(document_ids),
            Document.owner_id == owner_id,
            Document.owner_id.is_not(None),
        )
        .order_by(Document.created_at.asc(), Document.id.asc(), InvoiceLineItem.position.asc())
    )
    with archive.open("line_items.csv", mode="w") as binary_file:
        text_file = io.TextIOWrapper(binary_file, encoding="utf-8-sig", newline="")
        try:
            writer = csv.writer(text_file, lineterminator="\r\n")
            writer.writerow(LINE_ITEM_HEADERS)
            for row in database.execute(statement).yield_per(QUERY_BATCH_SIZE):
                writer.writerow(_line_item_row(*row))
            text_file.flush()
        finally:
            text_file.detach()


def _invoice_row(document: Document, extraction: InvoiceExtraction | None) -> tuple[str, ...]:
    if extraction is None:
        return (
            str(document.id),
            sanitize_csv_text(document.original_filename),
            document.status.value,
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
        )
    return (
        str(document.id),
        sanitize_csv_text(document.original_filename),
        document.status.value,
        sanitize_csv_text(extraction.invoice_number),
        serialize_date(extraction.invoice_date),
        sanitize_csv_text(extraction.vendor_name),
        sanitize_csv_text(extraction.customer_name),
        extraction.currency or "",
        serialize_money(extraction.subtotal),
        serialize_money(extraction.tax),
        serialize_money(extraction.total),
        "true" if extraction.is_valid else "false",
        sanitize_csv_text(extraction.validation_error),
    )


def _line_item_row(
    document_id: uuid.UUID,
    invoice_number: str | None,
    position: int,
    description: str,
    quantity: Decimal,
    unit_price: Decimal,
    line_total: Decimal,
) -> tuple[str, ...]:
    return (
        str(document_id),
        sanitize_csv_text(invoice_number),
        str(position),
        sanitize_csv_text(description),
        serialize_quantity(quantity),
        serialize_money(unit_price),
        serialize_money(line_total),
    )


# ---------------------------------------------------------------------------
# XLSX export
# ---------------------------------------------------------------------------

_INVOICE_COL_WIDTHS = (38, 30, 16, 20, 14, 25, 25, 10, 14, 14, 14, 10, 30)
_LINE_ITEM_COL_WIDTHS = (38, 20, 16, 40, 12, 12, 14)
_DATE_FORMAT = "YYYY-MM-DD"
_DATE_COL_INDEX = 5  # 1-based column index for invoice_date


def create_xlsx_export(
    database: Session, owner_id: uuid.UUID, requested_document_ids: Iterable[uuid.UUID]
) -> BinaryIO:
    document_ids = tuple(dict.fromkeys(requested_document_ids))
    if not document_ids:
        raise ValueError("At least one document ID is required.")

    _verify_owned_documents(database, owner_id, document_ids)
    export_file = tempfile.SpooledTemporaryFile(max_size=SPOOL_MEMORY_LIMIT, mode="w+b")
    try:
        workbook = Workbook()
        invoices_sheet = workbook.active
        assert invoices_sheet is not None
        invoices_sheet.title = "Invoices"
        line_items_sheet = workbook.create_sheet(title="Line Items")

        _write_invoices_sheet(invoices_sheet, database, owner_id, document_ids)
        _write_line_items_sheet(line_items_sheet, database, owner_id, document_ids)

        workbook.save(export_file)
        export_file.seek(0)
        return cast(BinaryIO, export_file)
    except Exception:
        export_file.close()
        raise


def _apply_sheet_formatting(
    sheet: Worksheet, headers: tuple[str, ...], col_widths: tuple[int, ...]
) -> None:
    for idx, width in enumerate(col_widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _write_invoices_sheet(
    sheet: Worksheet,
    database: Session,
    owner_id: uuid.UUID,
    document_ids: tuple[uuid.UUID, ...],
) -> None:
    sheet.append(list(INVOICE_HEADERS))
    _apply_sheet_formatting(sheet, INVOICE_HEADERS, _INVOICE_COL_WIDTHS)

    statement = (
        select(Document, InvoiceExtraction)
        .outerjoin(InvoiceExtraction, InvoiceExtraction.document_id == Document.id)
        .where(
            Document.id.in_(document_ids),
            Document.owner_id == owner_id,
            Document.owner_id.is_not(None),
        )
        .order_by(Document.created_at.asc(), Document.id.asc())
    )
    for document, extraction in database.execute(statement).yield_per(QUERY_BATCH_SIZE):
        sheet.append(_xlsx_invoice_row(document, extraction))
        date_cell = sheet.cell(row=sheet.max_row, column=_DATE_COL_INDEX)
        if date_cell.value is not None:
            date_cell.number_format = _DATE_FORMAT


def _write_line_items_sheet(
    sheet: Worksheet,
    database: Session,
    owner_id: uuid.UUID,
    document_ids: tuple[uuid.UUID, ...],
) -> None:
    sheet.append(list(LINE_ITEM_HEADERS))
    _apply_sheet_formatting(sheet, LINE_ITEM_HEADERS, _LINE_ITEM_COL_WIDTHS)

    statement = (
        select(
            Document.id,
            InvoiceExtraction.invoice_number,
            InvoiceLineItem.position,
            InvoiceLineItem.description,
            InvoiceLineItem.quantity,
            InvoiceLineItem.unit_price,
            InvoiceLineItem.line_total,
        )
        .select_from(Document)
        .join(InvoiceExtraction, InvoiceExtraction.document_id == Document.id)
        .join(
            InvoiceLineItem,
            InvoiceLineItem.invoice_document_id == InvoiceExtraction.document_id,
        )
        .where(
            Document.id.in_(document_ids),
            Document.owner_id == owner_id,
            Document.owner_id.is_not(None),
        )
        .order_by(Document.created_at.asc(), Document.id.asc(), InvoiceLineItem.position.asc())
    )
    for row in database.execute(statement).yield_per(QUERY_BATCH_SIZE):
        sheet.append(_xlsx_line_item_row(*row))


def _xlsx_invoice_row(document: Document, extraction: InvoiceExtraction | None) -> list[object]:
    if extraction is None:
        return [
            str(document.id),
            sanitize_cell_text(document.original_filename),
            document.status.value,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
        ]
    return [
        str(document.id),
        sanitize_cell_text(document.original_filename),
        document.status.value,
        sanitize_cell_text(extraction.invoice_number),
        extraction.invoice_date,
        sanitize_cell_text(extraction.vendor_name),
        sanitize_cell_text(extraction.customer_name),
        extraction.currency or None,
        _money_string(extraction.subtotal),
        _money_string(extraction.tax),
        _money_string(extraction.total),
        extraction.is_valid,
        sanitize_cell_text(extraction.validation_error) or None,
    ]


def _xlsx_line_item_row(
    document_id: uuid.UUID,
    invoice_number: str | None,
    position: int,
    description: str,
    quantity: Decimal,
    unit_price: Decimal,
    line_total: Decimal,
) -> list[object]:
    return [
        str(document_id),
        sanitize_cell_text(invoice_number) or None,
        position,
        sanitize_cell_text(description),
        _quantity_string(quantity),
        _money_string(unit_price),
        _money_string(line_total),
    ]


def _money_string(value: Decimal | None) -> str | None:
    return None if value is None else format(value, ".2f")


def _quantity_string(value: Decimal | None) -> str | None:
    return None if value is None else format(value, ".4f")


# ---------------------------------------------------------------------------
# JSON export
# ---------------------------------------------------------------------------


def create_json_export(
    database: Session, owner_id: uuid.UUID, requested_document_ids: Iterable[uuid.UUID]
) -> list[dict[str, Any]]:
    document_ids = tuple(dict.fromkeys(requested_document_ids))
    if not document_ids:
        raise ValueError("At least one document ID is required.")

    _verify_owned_documents(database, owner_id, document_ids)

    line_items_statement = (
        select(
            Document.id,
            InvoiceExtraction.invoice_number,
            InvoiceLineItem.position,
            InvoiceLineItem.description,
            InvoiceLineItem.quantity,
            InvoiceLineItem.unit_price,
            InvoiceLineItem.line_total,
        )
        .select_from(Document)
        .join(InvoiceExtraction, InvoiceExtraction.document_id == Document.id)
        .join(
            InvoiceLineItem,
            InvoiceLineItem.invoice_document_id == InvoiceExtraction.document_id,
        )
        .where(
            Document.id.in_(document_ids),
            Document.owner_id == owner_id,
            Document.owner_id.is_not(None),
        )
        .order_by(Document.created_at.asc(), Document.id.asc(), InvoiceLineItem.position.asc())
    )
    line_items_by_doc: dict[uuid.UUID, list[dict[str, Any]]] = {}
    for (
        doc_id,
        _inv_num,
        position,
        description,
        quantity,
        unit_price,
        line_total,
    ) in database.execute(line_items_statement).yield_per(QUERY_BATCH_SIZE):
        if doc_id not in line_items_by_doc:
            line_items_by_doc[doc_id] = []
        line_items_by_doc[doc_id].append(
            {
                "line_item_index": position,
                "description": description,
                "quantity": format(quantity, ".4f"),
                "unit_price": format(unit_price, ".2f"),
                "line_total": format(line_total, ".2f"),
            }
        )

    documents_statement = (
        select(Document, InvoiceExtraction)
        .outerjoin(InvoiceExtraction, InvoiceExtraction.document_id == Document.id)
        .where(
            Document.id.in_(document_ids),
            Document.owner_id == owner_id,
            Document.owner_id.is_not(None),
        )
        .order_by(Document.created_at.asc(), Document.id.asc())
    )

    results: list[dict[str, Any]] = []
    for document, extraction in database.execute(documents_statement).yield_per(QUERY_BATCH_SIZE):
        results.append(
            _json_document_entry(document, extraction, line_items_by_doc.get(document.id, []))
        )

    return results


def _json_document_entry(
    document: Document,
    extraction: InvoiceExtraction | None,
    line_items: list[dict[str, Any]],
) -> dict[str, Any]:
    if extraction is None:
        return {
            "document_id": str(document.id),
            "filename": document.original_filename,
            "document_status": document.status.value,
            "invoice_number": None,
            "invoice_date": None,
            "vendor_name": None,
            "customer_name": None,
            "currency": None,
            "subtotal": None,
            "tax": None,
            "total": None,
            "is_valid": None,
            "validation_error": None,
            "line_items": [],
        }

    return {
        "document_id": str(document.id),
        "filename": document.original_filename,
        "document_status": document.status.value,
        "invoice_number": extraction.invoice_number,
        "invoice_date": (
            extraction.invoice_date.isoformat() if extraction.invoice_date is not None else None
        ),
        "vendor_name": extraction.vendor_name,
        "customer_name": extraction.customer_name,
        "currency": extraction.currency,
        "subtotal": _money_string(extraction.subtotal),
        "tax": _money_string(extraction.tax),
        "total": _money_string(extraction.total),
        "is_valid": extraction.is_valid,
        "validation_error": extraction.validation_error,
        "line_items": line_items,
    }
